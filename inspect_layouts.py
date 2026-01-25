from pptx import Presentation
import os

ppt_path = r"D:\99T\Rasa\Project-IQ\presentations\GLPI_Chatbot_PPT.pptx"

if os.path.exists(ppt_path):
    print(f"Loading {ppt_path}...")
    try:
        prs = Presentation(ppt_path)
        print(f"Slide Width: {prs.slide_width}, Height: {prs.slide_height}")
        
        print("\n--- Slide Layouts in Master ---")
        for i, layout in enumerate(prs.slide_layouts):
            print(f"Layout Index {i}: {layout.name}")
            # print placeholders
            for shape in layout.placeholders:
                print(f"  - Placeholder idx {shape.placeholder_format.idx}, name: {shape.name}, type: {shape.placeholder_format.type}")
                
        print("\n--- Existing Slides ---")
        for i, slide in enumerate(prs.slides):
            layout = slide.slide_layout
            print(f"Slide {i+1} uses Layout {prs.slide_layouts.index(layout)}: {layout.name}")
            if slide.shapes.title:
                print(f"  Title: {slide.shapes.title.text}")
                
    except Exception as e:
        print(f"Error loading PPT: {e}")
else:
    print("File not found.")
